import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from styles import load_css
class AuditWorkPage:
   # read excel file
   AUDIT_FILE = "Audit.xlsx"
   def load_data(self):
       excel_file = pd.ExcelFile(self.AUDIT_FILE)
       records = []
       for sheet in excel_file.sheet_names:
           ws_df = pd.read_excel(self.AUDIT_FILE, sheet_name=sheet, header=None)
           for row in range(9, len(ws_df)):
               project_id = ws_df.iloc[row, 2]
               if pd.isna(project_id):
                   continue
               project_id = str(project_id).strip()
               if not project_id:
                   continue
               if project_id.lower().startswith("note"):
                   continue
               if "release" in project_id.lower():
                   continue
               if not project_id.startswith("CS"):
                   continue
               records.append({
                   "Sheet": sheet,
                   "Excel_Row": row + 1,  # 1-based indexing for openpyxl
                   "Project ID": project_id,
                   "Project Name": str(ws_df.iloc[row, 4]).strip(),
                   "Category": str(ws_df.iloc[row, 6]).strip(),
                   "Project Status": str(ws_df.iloc[row, 7]).strip(),
                   "Audit Type": str(ws_df.iloc[row, 8]).strip(),
                   "Audit Plan Date": str(ws_df.iloc[row, 9]).strip(),
                   "Auditee": str(ws_df.iloc[row, 10]).strip()
               })
       return pd.DataFrame(records)
   def save_record(self, sheet_name, excel_row, project_status, audit_type, audit_date, auditee):
       wb = load_workbook(self.AUDIT_FILE)
       ws = wb[sheet_name]
       # Mapping to match the 1-based column indices of your load_data function:
       # Index 7 -> Column 8 (H), Index 8 -> Column 9 (I), etc.
       ws.cell(row=excel_row, column=8, value=project_status)
       ws.cell(row=excel_row, column=9, value=audit_type)
       ws.cell(row=excel_row, column=10, value=audit_date)
       ws.cell(row=excel_row, column=11, value=auditee)
       wb.save(self.AUDIT_FILE)
   def show(self):
       # loading styles.py for styling the page.
       load_css()
       st.markdown("""
<div style='

        padding:20px;

        border-radius:12px;

        background:linear-gradient(90deg,#0f172a,#1e293b);

        color:white;

        margin-bottom:20px;

        '>
<h2 style='margin:0;'>Audit Work Dashboard</h2>
</div>
<style>

        .audit-table-card { border: 1px solid #7c5cff; border-radius: 16px; background: rgba(124,92,255,0.08); padding: 14px; margin-bottom: 24px; max-width: 100%; overflow-x: auto; }

        .audit-table { width: 100%; border-collapse: collapse; margin-top: 12px; font-size: 13px; }

        .audit-table th, .audit-table td { border: 1px solid rgba(124,92,255,0.16); padding: 8px 10px; text-align: left; }

        .audit-table th { background: rgba(124,92,255,0.12); color: #f8fbff; font-weight: 700; }

        .audit-table td { color: #eaf1ff; }

        .audit-table tr:hover { background: rgba(124,92,255,0.08); }

        /* --- TARGET ONLY THE FORM SUBMIT BUTTON (LEAVES SIDEBAR ALONE) --- */

        div[data-testid="stForm"] button[data-testid="stFormSubmitButton"] {

            background-color: #1e293b !important;

            color: #ffffff !important;

            border: 1px solid #7c5cff !important;

            border-radius: 8px !important;

            padding: 0.5rem 1rem !important;

            transition: all 0.3s ease !important;

        }

        /* --- HOVER STATE FOR FORM BUTTON ONLY --- */

        div[data-testid="stForm"] button[data-testid="stFormSubmitButton"]:hover {

            background-color: #7c5cff !important;

            color: #ffffff !important;

            border-color: #9075ff !important;

            box-shadow: 0 0 10px rgba(124, 92, 255, 0.5) !important;

        }
</style>
 
<style>

        .audit-table-card { border: 1px solid #7c5cff; border-radius: 16px; background: rgba(124,92,255,0.08); padding: 14px; margin-bottom: 24px; max-width: 100%; overflow-x: auto; }

        .audit-table { width: 100%; border-collapse: collapse; margin-top: 12px; font-size: 13px; }

        .audit-table th, .audit-table td { border: 1px solid rgba(124,92,255,0.16); padding: 8px 10px; text-align: left; }

        .audit-table th { background: rgba(124,92,255,0.12); color: #f8fbff; font-weight: 700; }

        .audit-table td { color: #eaf1ff; }

        .audit-table tr:hover { background: rgba(124,92,255,0.08); }

        /* --- TARGETS ONLY FORM SUBMIT BUTTONS GLOBALLY (LEAVES SIDEBAR ALONE) --- */

        button[kind="primaryFormSubmit"] {

            background-color: #1e293b !important;

            color: #ffffff !important;

            border: 1px solid #7c5cff !important;

            border-radius: 8px !important;

            padding: 0.5rem 1rem !important;

            transition: all 0.3s ease !important;

        }

        /* --- HOVER STATE --- */

        button[kind="primaryFormSubmit"]:hover {

            background-color: #7c5cff !important;

            color: #ffffff !important;

            border-color: #9075ff !important;

            box-shadow: 0 0 10px rgba(124, 92, 255, 0.5) !important;

        }
</style>
 
 

        """, unsafe_allow_html=True)
 
 
       try:
           df = self.load_data()
       except Exception as e:
           st.error(f"Error loading audit file: {e}")
           return
       # displaying metrics from the audit data
       c1, c2, c3 = st.columns(3)
       c1.metric("Total Audits", len(df))
       c2.metric("Projects", df["Project ID"].nunique())
       c3.metric("Weeks", df["Sheet"].nunique())
       st.divider()
       search = st.text_input("Search Project ID / Project Name")
       if search:
           df = df[
               df["Project ID"].str.contains(search, case=False, na=False) |
               df["Project Name"].str.contains(search, case=False, na=False)
           ]
       if df.empty:
           st.info("No matching records found.")
           return
       sheet_names = sorted(df["Sheet"].unique().tolist())
       # Different tabs for displaying various sheets
       tabs = st.tabs(sheet_names)
       role = st.session_state.get("role", "")
       for idx, sheet in enumerate(sheet_names):
           with tabs[idx]:
               sheet_df = df[df["Sheet"] == sheet]
               display_df = sheet_df.drop(columns=["Sheet", "Excel_Row"], errors="ignore").reset_index(drop=True)
               # HTML Table compilation
               table_html = '<div class="audit-table-card">'
               table_html += '<table class="audit-table">'
               table_html += '<thead><tr>'
               for col in display_df.columns:
                   table_html += f'<th>{col}</th>'
               table_html += '</tr></thead><tbody>'
               for _, row in display_df.iterrows():
                   table_html += '<tr>'
                   for value in row:
                       table_html += f'<td>{value}</td>'
                   table_html += '</tr>'
               table_html += '</tbody></table></div>'
               st.markdown(table_html, unsafe_allow_html=True)
               # Reporting Manager dashboard edit options
               if role == "Reporting Manager" and not sheet_df.empty:
                   st.subheader("Edit Audit")
                   selected_id = st.selectbox(
                       "Select Project",
                       sheet_df["Project ID"].tolist(),
                       key=f"project_{sheet}"
                   )
                   selected = sheet_df[sheet_df["Project ID"] == selected_id].iloc[0]
                   # Wrap editing elements inside a form context block to handle dynamic input updates cleanly
                   with st.form(key=f"edit_form_{sheet}"):
                       new_status = st.text_input(
                           "Project Status",
                           selected["Project Status"]
                       )
                       new_audit_type = st.text_input(
                           "Audit Type",
                           selected["Audit Type"]
                       )
                       new_date = st.text_input(
                           "Audit Plan Date",
                           selected["Audit Plan Date"]
                       )
                       new_auditee = st.text_input(
                           "Auditee",
                           selected["Auditee"]
                       )
                       submit_button = st.form_submit_button("Save Changes")
                       if submit_button:
                           self.save_record(
                               selected["Sheet"],
                               int(selected["Excel_Row"]),
                               new_status,
                               new_audit_type,
                               new_date,
                               new_auditee
                           )
                           st.success("Audit updated successfully.")
                           st.rerun()
def show_audit():
   # Make sure to mock a role for testing if needed, e.g.:
   # if "role" not in st.session_state: st.session_state["role"] = "Reporting Manager"
   AuditWorkPage().show()