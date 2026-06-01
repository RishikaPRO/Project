#team member->input data->store excel->data visualization
#fields->project id,project name, team members, progress
# 1) Enter project id, name
# 2) Display project id, name on top 

import streamlit as st
import pandas as pd
import altair as alt
import openpyxl
from openpyxl import Workbook, load_workbook
import os

st.set_page_config(page_title="Job Tracker", layout="wide")
st.title("Job Tracker Application")

filename="job_tracker.xlsx"
if not os.path.exists(filename):
    wb=Workbook()
    ws=wb.active
    ws.title="Job Tracker"
    ws.append(["Project ID", "Project Name", "Employee ID","Employee Name", "Job Undertaken", "Time Duration", "Progress"])
    #table header
    ws.add_table(displayName="JobTrackerTable", ref="A1:G1")
    wb.save(filename)

df = pd.read_excel(filename)
df.to_excel("job_tracker.xlsx", index=True)


if "project_selected" not in st.session_state:
    st.session_state.project_selected=False
if not st.session_state.project_selected:
   with st.form("project_form"):
       st.subheader("Select Project")
       project_id = st.text_input("Project ID")
       project_name = st.text_input("Project Name")
       start_button = st.form_submit_button("Open Project")
       if start_button:
           st.session_state.project_id = project_id
           st.session_state.project_name = project_name
           st.session_state.project_selected = True
           st.rerun()

else:
   project_id = st.session_state.project_id
   project_name = st.session_state.project_name
   st.markdown(f"# Project : {project_name}")
   st.markdown(f"### Project ID : {project_id}")
   st.divider()

   project_df = df[(df["Project ID"].astype(str) == str(project_id))]  #filter: Project ID

   st.subheader("Add Employee Job Details")
   with st.form("employee_form"):
       col1, col2 = st.columns(2)
       with col1:
           emp_id = st.text_input("Employee ID")
       with col2:
           emp_name = st.text_input("Employee Name")
       job = st.selectbox( "Job Undertaken", ["Requirement Document", "Design Document", "Coding", "ITP", "Testing", "Audit Work"] )
       duration = st.text_input("Time Duration")
       progress = st.slider( "Progress Percentage",0, 100, 0)
       save_btn = st.form_submit_button("Save Entry")
       if save_btn:
           new_row = {"Project ID": project_id, "Project Name": project_name, "Employee ID": emp_id, "Employee Name": emp_name, "Job Undertaken": job, "Time Duration": duration, "Progress": progress}  #(table for each emp)
           new_df = pd.DataFrame([new_row])
           updated_df = pd.concat([df, new_df],ignore_index=True)
           updated_df = updated_df.sort_values(by=["Employee Name"])
           updated_df.to_excel(filename, index=False)
           st.success("Data Saved Successfully!")
           st.rerun()
   st.divider()
   df.to_excel("Job_Tracker.xlsx", index=False)
   #reload
   df = pd.read_excel(filename)
   project_df = df[(df["Project ID"].astype(str) == str(project_id))]

   project_df = project_df.sort_values(by=["Employee Name"])

#table
   st.subheader("Employee Records")
   st.dataframe(project_df, use_container_width=True)
   st.divider()


   st.subheader("Employee Progress Tracking")
   employee_groups = project_df.groupby("Employee ID")
   for emp_id, emp_data in employee_groups:
       emp_name = emp_data.iloc[0]["Employee Name"]
       st.markdown(f"## Employee : {emp_name} ({emp_id})")
       for index, row in emp_data.iterrows():
           col1, col2 = st.columns([3, 3])
           with col1:
               st.write(f"### {row['Job Undertaken']}")
               st.write(f"Duration : {row['Time Duration']}")
           with col2:
               st.progress(int(row["Progress"]))
               st.write(f"{row['Progress']}% Completed")
           #with col3:  #audit in detail view button->audit details page
               #with st.expander("Audit View"):
                   #st.write(f"Project ID : {row['Project ID']}")
                   #st.write(f"Project Name : {row['Project Name']}")
                   #st.write(f"Employee ID : {row['Employee ID']}")
                   #st.write(f"Employee Name : {row['Employee Name']}")
                   #st.write(f"Job : {row['Job Undertaken']}")
                   #st.write(f"Time Duration : {row['Time Duration']}")
                   #st.write(f"Progress : {row['Progress']}%")
       st.divider()
   #graph 
   st.subheader("Progress Visualization")
   chart = alt.Chart(project_df).mark_bar().encode(
       x=alt.X("Employee Name:N",title="Employee"),
       y=alt.Y("Progress:Q",title="Progress %"),
       color="Job Undertaken:N",
       tooltip=["Employee Name","Employee ID","Job Undertaken","Progress"]
   ).properties( width=900, height=450)
   st.altair_chart(
       chart,
       use_container_width=True
   )
   st.divider()
   if st.button("Open Another Project"):
       st.session_state.project_selected = False
       st.rerun()
